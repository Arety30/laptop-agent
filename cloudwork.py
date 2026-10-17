"""Cloud-side work for the cloud bot: real files (PDF/DOCX/MD), images, and coding — all executed on
Google Cloud Shell via cloudshell.py, never on the laptop. The laptop only runs the small LLM that
writes the content / code.

    make_document(topic, fmt, out_dir)   -> {"ok", "file", "error"}
    make_image(prompt, out_dir)          -> {"ok", "file", "error"}
    code_task(task, out_dir)             -> {"ok", "out", "files", "code"}
"""
from __future__ import annotations

import os
import re
import time
import urllib.parse

import cloudshell
import llm

_DOC_SYS = ("You write complete, well-structured documents in Markdown. Use a # title, ## sections, "
            "bullet lists and tables where useful. Be accurate and specific. Output ONLY the markdown, "
            "no commentary, no code fences.")

_CODE_SYS = ("You write a single self-contained Python 3 script for Linux. Available: reportlab, fpdf2, "
             "markdown, weasyprint, pillow, matplotlib (use Agg backend), python-docx, openpyxl, python-pptx, "
             "numpy, pandas, requests. Save every output file in the CURRENT directory with a descriptive "
             "name, and print a short summary. Output ONLY the Python code inside one ```python block.")

# fixed, deterministic converter that runs on the VM (markdown -> pdf/docx/html)
_CONVERT = r'''
import sys, markdown, re
fmt, name = sys.argv[1], sys.argv[2]
md = open("doc.md", encoding="utf-8").read()
html = markdown.markdown(md, extensions=["tables", "fenced_code", "toc"])
css = """body{font-family:'DejaVu Sans',sans-serif;font-size:11pt;line-height:1.45;margin:2cm}
h1{color:#1a3c6e;border-bottom:2px solid #1a3c6e;padding-bottom:4px}h2{color:#2b5797;margin-top:1.2em}
table{border-collapse:collapse;width:100%;margin:8px 0}th,td{border:1px solid #999;padding:4px 8px;text-align:left}
th{background:#e8eef7}code{background:#f2f2f2;padding:1px 4px}pre{background:#f2f2f2;padding:8px;white-space:pre-wrap}"""
full = f"<html><head><meta charset='utf-8'><style>{css}</style></head><body>{html}</body></html>"
if fmt == "pdf":
    from weasyprint import HTML
    HTML(string=full).write_pdf(name + ".pdf")
elif fmt == "docx":
    from docx import Document
    d = Document()
    for line in md.splitlines():
        if line.startswith("# "): d.add_heading(line[2:], 0)
        elif line.startswith("## "): d.add_heading(line[3:], 1)
        elif line.startswith("### "): d.add_heading(line[4:], 2)
        elif re.match(r"^\s*[-*] ", line): d.add_paragraph(re.sub(r"^\s*[-*] ", "", line), style="List Bullet")
        elif re.match(r"^\s*\d+\. ", line): d.add_paragraph(re.sub(r"^\s*\d+\. ", "", line), style="List Number")
        elif line.strip(): d.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", line))
    d.save(name + ".docx")
elif fmt == "html":
    open(name + ".html", "w", encoding="utf-8").write(full)
else:
    open(name + ".md", "w", encoding="utf-8").write(md)
print("converted", fmt)
'''


def _slug(s: str, n: int = 40) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")
    return (s[:n] or "file").rstrip("_")


def _strip_fences(t: str) -> str:
    m = re.search(r"```(?:markdown|md)?\s*\n(.*?)```", t, re.S)
    return (m.group(1) if m else t).strip()


async def make_document(topic: str, fmt: str = "pdf", out_dir: str = "/tmp") -> dict:
    """LLM writes markdown on the laptop; the VM converts it to a real PDF/DOCX/HTML/MD file."""
    fmt = fmt if fmt in ("pdf", "docx", "html", "md") else "pdf"
    md = _strip_fences(await llm.complete("report", _DOC_SYS, f"Write a document about: {topic}"))
    if len(md) < 40:
        return {"ok": False, "error": "the model produced no usable content"}
    name = _slug(topic)
    stamp = int(time.time())
    remote_dir = f"jobs/{stamp}"
    r = await cloudshell.run(
        f"mkdir -p ~/{remote_dir} && cd ~/{remote_dir} && cat > doc.md <<'__MD__'\n{md}\n__MD__\n"
        f"cat > convert.py <<'__PY__'\n{_CONVERT}\n__PY__\n"
        f"python3 convert.py {fmt} {name} 2>&1 | tail -3", timeout=240)
    if not r["ok"] and "converted" not in r["out"]:
        return {"ok": False, "error": r["out"][-500:], "markdown": md}
    os.makedirs(out_dir, exist_ok=True)
    local = os.path.join(out_dir, f"{name}.{fmt}")
    if not await cloudshell.get(f"{remote_dir}/{name}.{fmt}", local):
        return {"ok": False, "error": "file was made on the VM but couldn't be fetched", "markdown": md}
    return {"ok": True, "file": local, "markdown": md}


async def make_image(prompt: str, out_dir: str = "/tmp", size: int = 768) -> dict:
    """Image generation (pollinations.ai, free, no key) — fetched by the VM, then pulled here."""
    q = urllib.parse.quote(prompt)
    seed = int(time.time()) % 100000
    name = f"{_slug(prompt, 30)}_{seed}.jpg"
    url = f"https://image.pollinations.ai/prompt/{q}?width={size}&height={size}&nologo=true&seed={seed}"
    r = await cloudshell.run(
        f"mkdir -p ~/images && cd ~/images && curl -sL --max-time 150 -o {name} '{url}' && file -b {name} | cut -c1-40",
        timeout=180)
    if not r["ok"] or "image" not in r["out"].lower():
        return {"ok": False, "error": r["out"][-300:] or "image service failed"}
    os.makedirs(out_dir, exist_ok=True)
    local = os.path.join(out_dir, name)
    if not await cloudshell.get(f"images/{name}", local):
        return {"ok": False, "error": "couldn't fetch the image from the VM"}
    return {"ok": True, "file": local}


async def code_task(task: str, out_dir: str = "/tmp", retries: int = 1) -> dict:
    """Coder model writes a script; it runs on the VM; on error the traceback is fed back once."""
    prompt = f"Task: {task}"
    code, last = "", {}
    for attempt in range(retries + 1):
        raw = await llm.complete("coder", _CODE_SYS, prompt)
        m = re.search(r"```(?:python|py)?\s*\n(.*?)```", raw, re.S)
        code = (m.group(1) if m else raw).strip()
        if not code:
            return {"ok": False, "out": "no code produced", "files": [], "code": ""}
        last = await cloudshell.run_script(code, "python", out_dir, timeout=240)
        if last["ok"]:
            break
        prompt = (f"Task: {task}\n\nYour previous script:\n```python\n{code}\n```\n"
                  f"failed with:\n{last['out'][-1200:]}\n\nFix it. Output ONLY the corrected full script.")
    last["code"] = code
    return last
